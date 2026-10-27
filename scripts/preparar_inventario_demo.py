"""Prepara el Excel autorizado sin conectarse a la API ni modificar el original."""

import argparse
from collections import Counter, defaultdict
import hashlib
import json
import os
from pathlib import Path
import random
import re
import sys
import unicodedata


LABS = ("L201", "L409", "L413", "L414")
ACTORS = ("marko", "aldo", "jason", "jorge")
STATES = {
    "OPERATIVO": "OPERATIVO",
    "INOPERATIVO": "INOPERATIVO",
    "NOOPERATIVO": "INOPERATIVO",
    "AVERIADO": "INOPERATIVO",
    "DANADO": "INOPERATIVO",
    "MALOGRADO": "INOPERATIVO",
    "MANTENIMIENTO": "MANTENIMIENTO",
}


def text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = "".join(c for c in str(value) if unicodedata.category(c) != "Cf").strip()
    return value or None


def key(value):
    value = unicodedata.normalize("NFKD", text(value) or "")
    return "".join(c for c in value.upper() if c.isalnum())


def optional(value):
    value = text(value)
    return None if key(value) in {"", "NA", "SN", "SINSERIE"} or value == "-" else value


def year(value):
    value = optional(value)
    if value is None:
        return None
    try:
        result = int(value)
    except ValueError as error:
        raise ValueError("El anio debe ser un entero o estar vacio") from error
    if not 1900 <= result <= 2100:
        raise ValueError("Anio fuera del rango 1900-2100")
    return result


def read_rows(source):
    import openpyxl

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    rows = []
    try:
        for lab in LABS:
            if lab not in workbook.sheetnames:
                raise ValueError(f"Falta la hoja {lab}")
            sheet = workbook[lab]
            header = next(sheet.iter_rows(min_row=1, max_row=1, max_col=14, values_only=True))
            expected = ("SERIEUTEC", "SERIE", "CODIGOINTERNO", "CATEGORIA", "SUBCATEGORIA", "NOMBRE", "MARCA", "MODELO", "ESTADO")
            if tuple(key(v) for v in header[:9]) != expected:
                raise ValueError(f"Encabezados inesperados en {lab}; no se asumira otra estructura")
            for index, values in enumerate(sheet.iter_rows(min_row=2, max_row=128 if lab == "L201" else None,
                                                         max_col=14, values_only=True), 2):
                # Filas con solo un estado rellenado no representan equipos.
                if not any(text(v) for v in values[:8]):
                    continue
                rows.append({"lab": lab, "row": index, "values": list(values)})
    finally:
        workbook.close()
    return rows


def code_base(row):
    original = optional(row["values"][2])
    if original is None:
        return f'{row["lab"]}-IMP-F{row["row"]:04d}'
    normalized = re.sub(r"\s*-\s*", "-", original.upper())
    normalized = re.sub(r"\s+", "-", normalized)
    return f'{row["lab"]}-{normalized}'


def build_plan(source):
    source = Path(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    rows = read_rows(source)
    labs = list(LABS)
    random.Random(int(digest[:16], 16)).shuffle(labs)
    assignments = dict(zip(ACTORS, labs))
    owners = {lab: user for user, lab in assignments.items()}
    bases = Counter(code_base(row) for row in rows)
    serials = defaultdict(list)
    for row in rows:
        serial = optional(row["values"][1])
        if serial:
            serials[serial.upper()].append((row["lab"], row["row"]))
    duplicates = {serial: locations for serial, locations in serials.items() if len(locations) > 1}
    equipment, pending = [], []
    codes = set()
    for row in rows:
        values, lab, index = row["values"], row["lab"], row["row"]
        code = code_base(row)
        if bases[code] > 1:
            code += f"-F{index:04d}"
        if len(code) > 50 or code in codes:
            raise ValueError(f"Codigo no valido o duplicado despues de normalizar: {lab}:{index}")
        codes.add(code)
        for column in (3, 4, 5):
            if not text(values[column]):
                raise ValueError(f"Falta categoria, subcategoria o nombre en {lab}:{index}")
        serial = optional(values[1])
        duplicate_serial = bool(serial and serial.upper() in duplicates)
        state = STATES.get(key(values[8]))
        original = {
            "codigo": text(values[2]), "serieUtec": text(values[0]),
            "serieFabricante": text(values[1]), "estado": text(values[8]),
            "ambiente": text(values[11]), "comentario": text(values[13]),
        }
        notes = ["Carga inicial de demostracion desde Excel.",
                 f"Origen Excel: {digest}; hoja={lab}; fila={index}",
                 "Valores originales: " + json.dumps(original, ensure_ascii=False)]
        if duplicate_serial:
            notes.append("Serie de fabricante repetida en el Excel: se conserva aqui y el campo queda vacio, pendiente de verificar.")
        if text(values[11]) and key(values[11]) != key(lab):
            notes.append(f"Laboratorio tomado de la hoja {lab} por indicacion del usuario.")
        payload = {
            "codigoInterno": code, "serieUtec": optional(values[0]),
            "numeroSerie": None if duplicate_serial else serial,
            "nombre": text(values[5]), "marca": optional(values[6]),
            "modelo": optional(values[7]), "estado": state, "anio": year(values[10]),
            "ordenCompra": optional(values[9]), "ubicacionInterna": optional(values[12]),
            "comentario": "\n".join(notes), "requiereMantenimiento": state == "MANTENIMIENTO",
            "idResponsable": None,
        }
        for field, maximum in {"nombre": 150, "serieUtec": 100, "numeroSerie": 100,
                               "marca": 100, "modelo": 100, "ordenCompra": 50,
                               "ubicacionInterna": 200}.items():
            if payload[field] is not None and len(payload[field]) > maximum:
                raise ValueError(f"{field} supera {maximum} caracteres en {lab}:{index}")
        for column in (3, 4):
            if len(text(values[column])) > 100:
                raise ValueError(f"Nombre de catalogo demasiado largo en {lab}:{index}")
        item = {"source_sheet": lab, "source_row": index, "lab": lab, "actor": owners[lab],
                "categoria": text(values[3]), "subcategoria": text(values[4]), "payload": payload}
        if state is None:
            item["reason"] = "Estado vacio o ambiguo: pendiente de confirmar; la API no admite estado null."
            pending.append(item)
        else:
            equipment.append(item)
    users = [{"userName": username, "nombre": username.capitalize(), "apellido": "Demo",
              "email": f"{username}@inventario.test", "rol": role, "activo": True,
              "cargo": "Cuenta de demostracion"}
             for username, role in (("jason", "ADMIN"), ("jorge", "GESTOR"), ("gabriela", "ADMIN"))]
    return {
        "version": 1, "source_sha256": digest, "source_name": source.name,
        "assignments": assignments, "users": users, "equipment": equipment, "pending": pending,
        "summary": {"candidates": len(rows), "ready": len(equipment), "pending": len(pending),
                    "ready_by_lab": dict(Counter(item["lab"] for item in equipment)),
                    "pending_by_lab": dict(Counter(item["lab"] for item in pending)),
                    "states": dict(Counter(item["payload"]["estado"] for item in equipment)),
                    "duplicated_manufacturer_serials": len(duplicates)},
    }


def write_plan(plan, directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / "plan.json"
    content = json.dumps(plan, ensure_ascii=False, indent=2) + "\n"
    # No reemplazar un plan distinto utilizado por una ejecucion anterior.
    if path.exists() and json.loads(path.read_text(encoding="utf-8")) != plan:
        raise ValueError("Ya existe un plan distinto en ese destino. Usa otra carpeta para evitar mezclar cargas.")
    temporary = path.with_suffix(".tmp")
    temporary.write_text(content, encoding="utf-8")
    os.replace(temporary, path)
    return path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archivo", type=Path, default=Path.home() / "Downloads" / "INVENTARIO Equipos grandes con código.xlsx")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    plan = build_plan(args.archivo)
    root = args.output or Path(os.environ.get("LOCALAPPDATA", Path.home())) / "inventario-laboratorios" / "importacion-demo" / plan["source_sha256"][:12]
    path = write_plan(plan, root)
    print(json.dumps({"plan": str(path), "assignments": plan["assignments"], **plan["summary"]}, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (ValueError, OSError) as error:
        print(f"Preparacion detenida: {error}", file=sys.stderr)
        sys.exit(1)
