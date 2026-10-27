"""Pruebas del preparador con libros sinteticos temporales, sin API ni BD."""

from collections import Counter
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl


MODULE_PATH = Path(__file__).with_name("preparar_inventario_demo.py")
SPEC = importlib.util.spec_from_file_location("preparar_inventario_demo", MODULE_PATH)
preparador = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(preparador)

HEADER = ["SERIE UTEC", "SERIE", "CÓDIGO INTERNO", "CATEGORÍA", "SUBCATEGORÍA",
          "NOMBRE", "MARCA", "MODELO", "ESTADO", "OC", "AÑO", "AMBIENTE",
          "UBICACIÓN", "Comentario"]


def record(lab, row=2, **changes):
    values = [f"UTEC-{lab}-{row}", f"SER-{lab}-{row}", f"EQ-{row:03d}",
              "Instrumentación", "Medición", f"Equipo {lab} {row}", "Marca", "Modelo",
              "Operativo", None, 2024, lab, "Mesa 1", None]
    for column, value in changes.items():
        values[HEADER.index(column)] = value
    return values


class PrepararInventarioTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="inventario-preparador-test-")
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / "inventario-sintetico.xlsx"

    def workbook(self, records=None, omit=(), wrong_header=None):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for lab in preparador.LABS:
            if lab in omit:
                continue
            sheet = workbook.create_sheet(lab)
            header = list(HEADER)
            if wrong_header == lab:
                header[2] = "Otro identificador"
            sheet.append(header)
            rows = records.get(lab, {}) if records is not None else {2: record(lab)}
            for index, values in rows.items():
                for column, value in enumerate(values, 1):
                    sheet.cell(index, column, value)
        workbook.save(self.source)
        workbook.close()
        return self.source

    @staticmethod
    def items(plan):
        return plan["equipment"] + plan["pending"]

    def test_l201_includes_128_excludes_129_and_skips_state_only_rows(self):
        records = {lab: {2: record(lab)} for lab in preparador.LABS}
        records["L201"].update({128: record("L201", 128), 129: record("L201", 129)})
        for index in range(3, 19):
            records["L409"][index] = [None] * 8 + ["OPERATIVO"] + [None] * 5
        self.workbook(records)
        rows = preparador.read_rows(self.source)
        self.assertEqual(Counter(row["lab"] for row in rows),
                         {"L201": 2, "L409": 1, "L413": 1, "L414": 1})
        self.assertIn(("L201", 128), [(row["lab"], row["row"]) for row in rows])
        self.assertNotIn(("L201", 129), [(row["lab"], row["row"]) for row in rows])

    def test_ambiguous_states_remain_null_and_pending_not_ready(self):
        records = {lab: {2: record(lab)} for lab in preparador.LABS}
        ambiguous = ["-", "Completo", "Incompleto", "En prestamo", None]
        for index in range(3, 30):
            records["L201"][index] = record("L201", index,
                                             **{"ESTADO": ambiguous[(index - 3) % 5]})
        for index in range(3, 7):
            records["L409"][index] = record("L409", index, **{"ESTADO": "Sin Batería"})
        plan = preparador.build_plan(self.workbook(records))
        self.assertEqual(plan["summary"]["pending"], 31)
        self.assertEqual(plan["summary"]["pending_by_lab"], {"L201": 27, "L409": 4})
        self.assertTrue(all(item["payload"]["estado"] is None for item in plan["pending"]))
        self.assertTrue(all(item["payload"]["estado"] == "OPERATIVO"
                            for item in plan["equipment"]))
        self.assertTrue(all("API no admite estado null" in item["reason"]
                            for item in plan["pending"]))

    def test_meaningful_state_aliases_preserve_operational_meaning(self):
        aliases = {"operativo": "OPERATIVO", "No Operativo": "INOPERATIVO",
                   "Averiado": "INOPERATIVO", "DAÑADO": "INOPERATIVO",
                   "MALOGRADO": "INOPERATIVO", "MANTENIMIENTO": "MANTENIMIENTO"}
        records = {lab: {} for lab in preparador.LABS}
        for index, state in enumerate(aliases, 2):
            records["L413"][index] = record("L413", index, **{"ESTADO": state})
        plan = preparador.build_plan(self.workbook(records))
        self.assertEqual(len(plan["pending"]), 0)
        for item in plan["equipment"]:
            original = json.loads(item["payload"]["comentario"].split("Valores originales: ")[1])
            self.assertEqual(item["payload"]["estado"], aliases[original["estado"]])
            self.assertEqual(item["payload"]["requiereMantenimiento"],
                             item["payload"]["estado"] == "MANTENIMIENTO")

    def test_duplicate_manufacturer_serial_is_null_on_both_sides_with_original_notes(self):
        records = {lab: {2: record(lab)} for lab in preparador.LABS}
        records["L201"][2][1] = "31F4C5C"
        records["L201"][3] = record("L201", 3, **{"SERIE": "31f4c5c"})
        records["L409"][2][1] = "C030148"
        records["L414"][2][1] = "C030148"
        records["L409"][3] = record("L409", 3, **{"SERIE": "C030145"})
        records["L414"][3] = record("L414", 3, **{"SERIE": "\u200bC030145"})
        plan = preparador.build_plan(self.workbook(records))
        self.assertEqual(plan["summary"]["duplicated_manufacturer_serials"], 3)
        affected = [item for item in plan["equipment"] if item["lab"] != "L413"]
        self.assertEqual(len(affected), 6)
        for item in affected:
            self.assertIsNone(item["payload"]["numeroSerie"])
            self.assertIn('"serieFabricante": "', item["payload"]["comentario"])
            self.assertIn("Serie de fabricante repetida", item["payload"]["comentario"])

    def test_codes_are_stable_unique_prefixed_and_missing_codes_use_source_row(self):
        records = {lab: {2: record(lab)} for lab in preparador.LABS}
        records["L201"][2][2] = "INS - OSC - 009"
        records["L201"][124] = record("L201", 124, **{"CÓDIGO INTERNO": "INS-OSC-009"})
        records["L201"][125] = record("L201", 125, **{"CÓDIGO INTERNO": None})
        records["L201"][126] = record("L201", 126, **{"CÓDIGO INTERNO": "-"})
        source = self.workbook(records)
        first, second = preparador.build_plan(source), preparador.build_plan(source)
        self.assertEqual(first, second)
        items = self.items(first)
        codes = [item["payload"]["codigoInterno"] for item in items]
        self.assertEqual(len(codes), len(set(codes)))
        self.assertTrue(all(code.startswith(item["lab"] + "-") and len(code) <= 50
                            for item, code in zip(items, codes)))
        self.assertIn("L201-INS-OSC-009-F0002", codes)
        self.assertIn("L201-INS-OSC-009-F0124", codes)
        self.assertIn("L201-IMP-F0125", codes)
        self.assertIn("L201-IMP-F0126", codes)

    def test_row_126_uses_l201_and_preserves_original_l209_reference(self):
        records = {lab: {2: record(lab)} for lab in preparador.LABS}
        records["L201"][126] = record("L201", 126,
                                          **{"AMBIENTE": "L209", "CÓDIGO INTERNO": None})
        plan = preparador.build_plan(self.workbook(records))
        item = next(item for item in plan["equipment"]
                    if item["lab"] == "L201" and item["source_row"] == 126)
        self.assertEqual(item["source_sheet"], "L201")
        self.assertIn('"ambiente": "L209"', item["payload"]["comentario"])
        self.assertIn("Laboratorio tomado de la hoja L201", item["payload"]["comentario"])

    def test_identifiers_preserve_leading_zeroes_and_remove_invisible_format_characters(self):
        records = {lab: {2: record(lab)} for lab in preparador.LABS}
        records["L201"][2][0] = "002079"
        records["L201"][2][1] = "\u200b802196010717810330"
        records["L409"][2][0] = 12345.0
        records["L409"][2][1] = "-"
        records["L413"][2][0] = "25073 - 25074 - 25075"
        records["L413"][2][1] = "28007 - 28006 - 28167"
        records["L414"][2][10] = "-"
        plan = preparador.build_plan(self.workbook(records))
        by_lab = {item["lab"]: item["payload"] for item in plan["equipment"]}
        self.assertEqual(by_lab["L201"]["serieUtec"], "002079")
        self.assertEqual(by_lab["L201"]["numeroSerie"], "802196010717810330")
        self.assertEqual(by_lab["L409"]["serieUtec"], "12345")
        self.assertIsNone(by_lab["L409"]["numeroSerie"])
        self.assertEqual(by_lab["L413"]["serieUtec"], "25073 - 25074 - 25075")
        self.assertEqual(by_lab["L413"]["numeroSerie"], "28007 - 28006 - 28167")
        self.assertIsNone(by_lab["L414"]["anio"])

    def test_assignments_are_reproducible_one_lab_per_actor_without_gabriela(self):
        source = self.workbook()
        first, second = preparador.build_plan(source), preparador.build_plan(source)
        assignments = first["assignments"]
        self.assertEqual(assignments, second["assignments"])
        self.assertEqual(set(assignments), {"marko", "aldo", "jason", "jorge"})
        self.assertEqual(set(assignments.values()), set(preparador.LABS))
        self.assertNotIn("gabriela", assignments)
        for item in self.items(first):
            self.assertEqual(assignments[item["actor"]], item["lab"])
        self.assertEqual({user["userName"]: user["rol"] for user in first["users"]},
                         {"jason": "ADMIN", "jorge": "GESTOR", "gabriela": "ADMIN"})
        self.assertTrue(all(user["email"].endswith("@inventario.test") for user in first["users"]))
        self.assertTrue(all("password" not in user for user in first["users"]))

    def test_source_bytes_and_sha_are_unchanged_after_preparation(self):
        source = self.workbook()
        before = source.read_bytes()
        plan = preparador.build_plan(source)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(plan["source_sha256"], hashlib.sha256(before).hexdigest())

    def test_headers_and_missing_sheets_fail_instead_of_guessing_columns(self):
        with self.subTest("unexpected_header"):
            with self.assertRaisesRegex(ValueError, "Encabezados inesperados en L413"):
                preparador.build_plan(self.workbook(wrong_header="L413"))
        with self.subTest("missing_sheet"):
            with self.assertRaisesRegex(ValueError, "Falta la hoja L414"):
                preparador.build_plan(self.workbook(omit=("L414",)))

    def test_invalid_code_length_and_missing_business_fields_fail_preparation(self):
        for changes, message in [({"CÓDIGO INTERNO": "X" * 47}, "Codigo no valido"),
                                 ({"NOMBRE": None}, "Falta categoria, subcategoria o nombre"),
                                 ({"AÑO": 1800}, "Anio fuera del rango")]:
            with self.subTest(changes=changes):
                records = {lab: {2: record(lab)} for lab in preparador.LABS}
                records["L201"][2] = record("L201", **changes)
                with self.assertRaisesRegex(ValueError, message):
                    preparador.build_plan(self.workbook(records))

    def test_saved_plan_is_idempotent_and_rejects_different_source(self):
        plan = preparador.build_plan(self.workbook())
        output = Path(self.directory.name) / "output"
        path = preparador.write_plan(plan, output)
        before = path.read_bytes()
        self.assertEqual(preparador.write_plan(plan, output), path)
        self.assertEqual(path.read_bytes(), before)
        changed = dict(plan, source_sha256="different-source")
        with self.assertRaisesRegex(ValueError, "Ya existe un plan distinto"):
            preparador.write_plan(changed, output)
        self.assertEqual(path.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
